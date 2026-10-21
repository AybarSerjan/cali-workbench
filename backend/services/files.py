from __future__ import annotations

import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

from pypdf import PdfReader
from pptx import Presentation
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from backend.config import ROOT
from backend.models import CalibrationSession, SourceFile, WikiRevision


# Formats MarkItDown converts to Markdown for us (plus the ones we still have a
# native extractor for). Office/PDF/eBook/structured-text/archive — no image OCR
# or audio, those need cloud/online backends we don't ship (see plan).
KNOWLEDGE_SUFFIXES = {
    ".pdf", ".md", ".markdown", ".json", ".pptx",
    ".docx", ".xlsx", ".xls", ".csv",
    ".html", ".htm", ".xml", ".epub", ".zip", ".msg", ".txt",
}
DATA_SUFFIXES = {".mdf", ".mf4", ".dat", ".a2l", ".hex"}


def sha256_stream(stream: BinaryIO) -> tuple[str, int, bytes]:
    digest = hashlib.sha256()
    chunks: list[bytes] = []
    size = 0
    while chunk := stream.read(1024 * 1024):
        digest.update(chunk)
        chunks.append(chunk)
        size += len(chunk)
    return digest.hexdigest(), size, b"".join(chunks)


def safe_filename(name: str) -> str:
    value = Path(name).name
    value = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "-", value).strip(".-")
    return value or "file"


def _persist_source(session: Session, filename: str, suffix: str, media_type: str, content: bytes, source_url: str = "") -> SourceFile:
    digest = hashlib.sha256(content).hexdigest()
    existing = session.scalar(select(SourceFile).where(SourceFile.sha256 == digest))
    if existing:
        return existing
    relative = Path("raw") / "files" / digest[:2] / f"{digest}-{safe_filename(filename)}"
    destination = ROOT / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(content)
    source = SourceFile(
        original_name=filename,
        media_type=media_type or "application/octet-stream",
        suffix=suffix,
        sha256=digest,
        relative_path=relative.as_posix(),
        size=len(content),
        source_url=source_url,
    )
    session.add(source)
    session.commit()
    session.refresh(source)
    return source


def store_source(session: Session, stream: BinaryIO, filename: str, media_type: str) -> SourceFile:
    suffix = Path(filename).suffix.lower()
    if suffix not in KNOWLEDGE_SUFFIXES | DATA_SUFFIXES:
        raise ValueError(f"不支持的知识文件类型: {suffix or '(无扩展名)'}")
    _, _, content = sha256_stream(stream)
    return _persist_source(session, filename, suffix, media_type or "application/octet-stream", content)


def store_link_source(session: Session, url: str, title: str, text: str) -> SourceFile:
    """External link knowledge (e.g. NotebookLM): the pasted text becomes a
    content-addressed .md snapshot, the URL is kept as provenance. The user is the
    trust boundary — we store what they pasted, we never fetch or execute it."""
    url = url.strip()
    if not re.match(r"^https?://", url, re.IGNORECASE):
        raise ValueError("链接需以 http:// 或 https:// 开头")
    title = title.strip()
    if not title:
        raise ValueError("标题不能为空")
    if not text.strip():
        raise ValueError("粘贴正文不能为空")
    filename = f"{safe_filename(title)}.md"
    return _persist_source(session, filename, ".md", "text/markdown", text.encode("utf-8"), source_url=url)


def delete_source(session: Session, source: SourceFile) -> None:
    wiki_refs = session.scalar(select(func.count()).select_from(WikiRevision).where(WikiRevision.source_id == source.id)) or 0
    calibration_refs = session.scalar(select(func.count()).select_from(CalibrationSession).where(CalibrationSession.a2l_file_id == source.id)) or 0
    if wiki_refs or calibration_refs:
        raise ValueError("知识源文件正在被 Wiki 或标定会话引用，无法删除")
    path = ROOT / source.relative_path
    session.delete(source)
    session.commit()
    path.unlink(missing_ok=True)


@dataclass
class Evidence:
    locator: str
    text: str


def extract_source(source: SourceFile) -> list[Evidence]:
    path = ROOT / source.relative_path
    suffix = source.suffix
    if suffix in {".md", ".markdown"}:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        return [Evidence(f"line {i}", line.strip()) for i, line in enumerate(lines, 1) if line.strip()]
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        output: list[Evidence] = []

        def walk(value, pointer=""):
            if isinstance(value, dict):
                for key, item in value.items():
                    walk(item, f"{pointer}/{str(key).replace('~', '~0').replace('/', '~1')}")
            elif isinstance(value, list):
                for index, item in enumerate(value):
                    walk(item, f"{pointer}/{index}")
            else:
                output.append(Evidence(pointer or "/", str(value)))

        walk(data)
        return output
    if suffix == ".pdf":
        output = []
        for number, page in enumerate(PdfReader(str(path)).pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                output.append(Evidence(f"page {number}", text))
        return output
    if suffix == ".pptx":
        output = []
        for number, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                output.append(Evidence(f"slide {number}", "\n".join(texts)))
        return output
    raise ValueError(f"缺少抽取器: {suffix}")


_markitdown = None


def _md_converter():
    """Lazily build a single MarkItDown instance. Kept local so importing this
    module never hard-fails if the optional extra is missing at runtime."""
    global _markitdown
    if _markitdown is None:
        from markitdown import MarkItDown  # imported lazily; heavy optional dep

        _markitdown = MarkItDown()
    return _markitdown


def source_markdown(source: SourceFile) -> str:
    """Convert a stored source to Markdown via MarkItDown. Falls back to the
    native evidence extractors (so PDF/PPTX/MD/JSON still work) if MarkItDown is
    unavailable or errors on the file. Content is untrusted — we only transform
    it, never execute anything inside it."""
    path = ROOT / source.relative_path
    try:
        text = (_md_converter().convert(str(path)).text_content or "").strip()
        if text:
            return text
    except Exception:
        pass  # fall through to the deterministic extractors below
    try:
        return "\n\n".join(f"**{item.locator}:** {item.text}" for item in extract_source(source)).strip()
    except ValueError:
        return ""


def copy_immutable(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        shutil.copy2(source, destination)

